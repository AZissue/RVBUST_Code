# -*- coding: utf-8 -*-
"""生成《3D 相机拼接方案》市场宣传文案 Word 文档。"""
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(BASE, "..", "3D相机拼接方案宣传文案.docx")

CN_FONT = "微软雅黑"
INK = (0x1E, 0x29, 0x3B)
BLUE = (0x15, 0x63, 0xEB)
TEAL = (0x0D, 0x94, 0x88)
AMBER = (0xB4, 0x5E, 0x09)
SUB = (0x64, 0x74, 0x8B)


def set_font(run, size=10.5, bold=False, color=None):
    run.font.name = CN_FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)


def para(doc, text, size=10.5, bold=False, align=None, space_after=6,
         color=None, indent=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_font(run, size=size, bold=bold, color=color)
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.35
    if indent:
        p.paragraph_format.first_line_indent = Pt(21)
    return p


def rich(doc, segments, size=10.5, space_after=6, indent=True):
    p = doc.add_paragraph()
    for text, bold, color in segments:
        run = p.add_run(text)
        set_font(run, size=size, bold=bold, color=color)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.35
    if indent:
        p.paragraph_format.first_line_indent = Pt(21)
    return p


def bullet(doc, text, bold_head=None, color=INK):
    p = doc.add_paragraph(style="List Bullet")
    if bold_head:
        run = p.add_run(bold_head)
        set_font(run, size=10.5, bold=True, color=color)
    run = p.add_run(text)
    set_font(run, size=10.5)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.3
    return p


def heading(doc, text, color=INK, size=14, space_before=10):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_font(run, size=size, bold=True, color=color)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(6)
    return p


def set_cell(cell, text, bold=False, size=9.5, shade=None, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    set_font(run, size=size, bold=bold, color=color)
    p.paragraph_format.space_after = Pt(2)
    if shade:
        tcPr = cell._tc.get_or_add_tcPr()
        shd = tcPr.makeelement(qn("w:shd"), {
            qn("w:val"): "clear", qn("w:fill"): shade})
        tcPr.append(shd)


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.4)
    sec.top_margin, sec.bottom_margin = Cm(2.4), Cm(2.4)

    # 主标题区
    para(doc, "把“拍不全”的工件，拼成一个完整的 3D", size=21, bold=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8, color=INK)
    para(doc, "两大类拼接方案 · 大件全覆盖，小件无死角", size=12,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4, color=SUB)
    para(doc, "单台 3D 相机一次只能拍一个视角：大件一张装不下，形状不规则的小件有遮挡死角。"
              "我们的拼接方案把多视角照片自动拼成一个完整 3D 模型——无论是大型工件的全表面"
              "采集，还是小型不规则零件的逆向建模，全程无需人工对齐。",
         size=10.5, indent=True, space_after=14)

    # 方案一
    heading(doc, "方案一：标记物拼接 —— 贴几个标记，剩下的交给软件", BLUE)
    rich(doc, [("一句话原理：", True, BLUE),
               ("在工件表面贴上带编号的编码标记，相机自动认出“两张照片里的同一个标记”，"
                "据此算出两片数据的相对位置，逐帧拼合成整体。", False, INK)])
    rich(doc, [("适合谁：", True, INK)])
    bullet(doc, "现场没有导轨、机器人、转台，也不想改造产线的客户；")
    bullet(doc, "筒体、模具、铸件、叶片、框架类大型工件——搬不动、转不了、拍不全；")
    bullet(doc, "形状不规则的小工件——凹槽、倒扣、深腔被自身遮挡，单视角拍不透，"
                "需要多视角拼接做逆向建模；")
    bullet(doc, "一次性扫描、逆向建模、大件测绘类项目。")
    rich(doc, [("解决什么痛点：", True, INK)])
    bullet(doc, "：大工件单视角装不下、小工件单视角有死角 → 多视角拍、自动拼，工件不用动；",
           bold_head="全")
    bullet(doc, "：不想买运动机构、不想改产线 → 零工装投入，贴标就能干；", bold_head="省")
    bullet(doc, "：人工拼接慢、拼不准 → 全自动，典型精度 0.5~2 mm。", bold_head="快")

    # 方案二
    heading(doc, "方案二：已知变换拼接 —— 设备知道相机在哪，拍下来就是拼好的", TEAL)
    rich(doc, [("一句话原理：", True, TEAL),
               ("相机装在导轨、机器人或转台上（或采用多相机固定阵列），每张照片拍摄时的位置"
                "关系由设备精确给出、一次标定长期有效，点云自动归位，无需逐帧计算。", False, INK)])
    rich(doc, [("适合谁：", True, INK)])
    bullet(doc, "产线在线检测、批量生产 → 多相机固定阵列，拍照即出结果；")
    bullet(doc, "圆柱、回转体工件 360° 全周检测 → 转台/托辊方案；")
    bullet(doc, "长筒、管道类长工件 → 编码器线扫，走一趟扫一条；")
    bullet(doc, "多品种、柔性换产的自动化产线 → 机器人 + 手眼标定。")
    rich(doc, [("解决什么痛点：", True, INK)])
    bullet(doc, "：外观件不能贴标记 → 工件上什么都不用贴；", bold_head="净")
    bullet(doc, "：节拍要求高、要全自动 → 拍照即拼接，无人工环节；", bold_head="快")
    bullet(doc, "：精度要求高且长期稳定 → 典型精度 0.1~1 mm，标定一次长期复用。", bold_head="精")

    # 对比表
    heading(doc, "一张表看懂怎么选", AMBER)
    t = doc.add_table(rows=5, cols=2)
    t.style = "Table Grid"
    rows = [
        ("你的现场情况", "推荐方案"),
        ("无运动机构、允许贴标、大件/一次性扫描", "标记物拼接"),
        ("产线在线检测、批量工件", "多相机固定阵列"),
        ("圆柱回转体、工件在滚轮架上", "转台 / 托辊"),
        ("长筒管道等长条形工件 / 多品种柔性产线", "编码器线扫 / 机器人 + 手眼标定"),
    ]
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            set_cell(t.rows[r].cells[c], val, bold=(r == 0),
                     shade="D9E2F3" if r == 0 else None)
    doc.add_paragraph()

    # 一句话总结
    para(doc, "一句话总结：有没有运动机构决定走哪条路——有机构走“已知变换”"
              "（快、精、全自动），没机构走“标记物”（零改造、最灵活）。两类方案覆盖从"
              "大件离线测绘到产线在线检测的全部场景。",
         size=11, bold=True, indent=True, space_after=4, color=INK)
    para(doc, "（典型精度为工程经验区间，具体以现场标定与验收测试为准。）",
         size=8.5, color=SUB, align=WD_ALIGN_PARAGRAPH.RIGHT)

    doc.save(OUT)
    print("saved:", os.path.abspath(OUT))


if __name__ == "__main__":
    build()
